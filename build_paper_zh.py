"""build_paper_zh.py — assemble the Chinese Research Article .docx."""
import os
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
RES = os.path.join(HERE, "..", "results")
OUT = os.path.join(HERE, "..", "Research_Article_ThinFilm_MLP_zh.docx")

doc = Document()

style = doc.styles["Normal"]
style.font.name = "Times New Roman"
style.font.size = Pt(10.5)
style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

for s in doc.sections:
    s.top_margin = Cm(2.2); s.bottom_margin = Cm(2.2)
    s.left_margin = Cm(2.4); s.right_margin = Cm(2.4)


def set_cn(run, size=10.5, bold=False, italic=False):
    run.font.size = Pt(size); run.bold = bold; run.italic = italic
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")


def h(text, size=13, space_before=10):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(4)
    set_cn(p.add_run(text), size=size, bold=True)
    return p


def para(text, size=10.5, italic=False, align=None, after=6):
    p = doc.add_paragraph()
    if align: p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.2
    set_cn(p.add_run(text), size=size, italic=italic)
    return p


def fig(fname, caption, width=14.5):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(os.path.join(RES, fname), width=Cm(width))
    c = doc.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_cn(c.add_run(caption), size=9, italic=True)
    c.paragraph_format.space_after = Pt(8)


# ---------------- 标题 ----------------
t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_cn(t.add_run("基于多层感知机的多层介质薄膜光谱预测与辅助设计"), size=16, bold=True)
t2 = doc.add_paragraph(); t2.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_cn(t2.add_run("MLP-Based Spectral Prediction and Data-Driven Design of Multilayer Dielectric Thin Films"),
       size=11)
t2.paragraph_format.space_after = Pt(8)

info = doc.add_paragraph(); info.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_cn(info.add_run("姓名：__________    学号：__________\n"), size=10.5)
set_cn(info.add_run("λtarget：480 nm    seed：270034    design_seed：270035"), size=10.5)

# ---------------- 摘要 ----------------
h("摘要", 12, space_before=8)
para(
    "多层介质薄膜的设计需要反复计算光谱，而每一次结构试算通常都要走一遍传输矩阵法"
    "（TMM），穷举筛选代价很高。本文把这个问题做成一个小型“科学智能”流程：用 TMM "
    "为 Air/H/L/H/L/Glass 膜系（nH=2.30，nL=1.45，ns=1.52）生成 5000 组“膜厚—反射率”"
    "数据，再训练一个 4–128–128–64–41 的多层感知机（MLP），把四层膜厚（40–180 nm）映射"
    "到 400–800 nm 内 41 个波长点的反射率。在留出测试集上，网络预测的 RMSE 为 0.0060；"
    "随着训练样本从 500 增加到 4000，误差单调下降（0.026→0.006）且边际收益递减。使用"
    "独立的 design_seed，MLP 可在毫秒级筛选 10000 个新候选，再经精确 TMM 复核，得到 "
    "480 nm 处反射率仅 0.0003 的减反设计。MLP 是准确而廉价的代理模型，TMM 仍是最终的"
    "物理校验。")
kw = doc.add_paragraph()
set_cn(kw.add_run("关键词："), bold=True)
set_cn(kw.add_run("光学薄膜；传输矩阵法；多层感知机；代理模型；科学智能"))

# ---------------- 1 引言 ----------------
h("1. 引言")
para(
    "多层介质薄膜由若干层透明薄层堆叠而成，其光谱响应来自各界面反射光之间的干涉，由"
    "层数、每层材料折射率和物理厚度共同决定，因此是设计变量的高度非线性函数。这类膜系"
    "广泛用作增透膜、高反膜、带通滤光片和辐射制冷涂层。")
para(
    "在材料固定的前提下，由膜厚正算反射谱是经典问题，可由传输矩阵法精确求解——它用一个 "
    "2×2 特征矩阵逐层推导电场与磁场。单次 TMM 很快，但即便只有四层，连续膜厚空间的搜索"
    "仍需要成千上万次重复计算，而“由目标光谱反求膜厚”的逆问题没有解析解。")
para(
    "机器学习提供了另一条路：用有限次 TMM 仿真训练一个代理模型，学习膜厚到光谱的非线性"
    "映射，之后预测新结构可以快几个数量级。把它用于粗筛大量候选，再只对少数最优结构做"
    "精确 TMM，就是典型的“物理+数据”科学智能流程。")
para(
    "本文把问题控制得小而可复现，回答三个问题：Q1，简单 MLP 能否从四个膜厚准确预测整条"
    "反射谱？Q2，训练数据增加时误差如何变化？Q3，MLP 能否作为代理快速筛选目标波长设计，"
    "其结果又如何与 TMM 对齐？")
fig("fig1_workflow.png", "图 1. 总体流程：TMM 正仿真 → 数据集 → MLP 代理模型 → 10000 候选快速筛选 → TMM 复核最优设计。",
    width=15.5)

# ---------------- 2 材料与方法 ----------------
h("2. 材料与方法")
h("2.1 光学模型与传输矩阵法", 11, space_before=6)
para(
    "膜系为 Air / H / L / H / L / Glass，共四层功能膜，正入射；忽略吸收与色散，正入射下 "
    "s、p 偏振简并。折射率取 n0=1.00（空气）、nH=2.30、nL=1.45、ns=1.52（玻璃基底）。"
    "第 j 层相位厚度 δj = 2π·nj·dj/λ，其特征矩阵为")
para("    Mj = [ [ cos δj,   i·sinδj / nj ],", italic=True, after=0)
para("          [ i·nj·sinδj,   cosδj      ] ]", italic=True, after=6)
para(
    "总矩阵 M = M1·M2·M3·M4。记 M = [[m11,m12],[m21,m22]]，由基底导纳得 B = m11 + "
    "m12·ns，C = m21 + m22·ns；反射系数 r = (B − C)/(B + C)（空气导纳取 1），反射率 "
    "R = |r|²。代码用裸玻璃极限 R = ((ns−1)/(ns+1))² ≈ 0.0426 做了自检，结果一致。")
fig("fig2_stack.png", "图 2. 物理模型：Air/H/L/H/L/Glass 四层膜系，正入射下的入射光 I(λ) 与反射光 R(λ)。",
    width=8.5)

h("2.2 个性化目标波长与随机种子", 11, space_before=6)
para(
    "学号后两位 N = 34，故 λtarget = 450 + 10×(N mod 31) = 480 nm。主种子 seed = 270034，"
    "控制数据生成、固定划分和 MLP 初始化（random、numpy、torch 三处均设为 270034）。为避免"
    "训练数据与筛选候选重复，设计阶段使用 design_seed = seed + 1 = 270035。")

h("2.3 数据集生成", 11, space_before=6)
para(
    "以 seed = 270034 在每层 [40, 180] nm 内均匀采样 5000 组膜厚向量，用 TMM 在 400–800 nm、"
    "步长 10 nm（共 41 点）上计算反射率，再做一次固定随机排列，划分为 4000 训练 / 500 验证 / "
    "500 测试。该划分在所有实验中保持不变。")

h("2.4 MLP 代理模型", 11, space_before=6)
para(
    "网络输入为标准化后的四个膜厚 [d1,d2,d3,d4]，输出为 41 个反射率，结构固定为 "
    "4→128→128→64→41（隐藏层 ReLU，输出层线性）。输入用训练集均值方差标准化，反射率目标"
    "保持在 [0,1]。损失为均方误差，Adam 优化（学习率 1e-3，batch 64），验证集 200 个 epoch "
    "无改善则早停。")
fig("fig3_mlp.png", "图 3. MLP 结构：4 个膜厚输入 → 两个 128 单元隐藏层 → 64 单元层 → 41 个反射率输出。",
    width=14)

h("2.5 训练数据量实验", 11, space_before=6)
para(
    "为研究 Q2，在固定训练池中依次取前 500、1000、2000、4000 个样本重训同一结构，验证集和"
    "测试集始终为固定的 500 个样本，其余优化设置保持不变。")

h("2.6 MLP 辅助薄膜筛选", 11, space_before=6)
para(
    "用 design_seed = 270035 生成 10000 个未参与训练的新候选，训练好的 MLP 预测其光谱，按 "
    "R(480 nm) 升序排列（减反目标），取前 10 名用 TMM 精确重算，最后按 TMM 结果重新排序取 "
    "Top5。")

# ---------------- 3 结果 ----------------
h("3. 结果")
h("3.1 MLP 训练与光谱预测", 11, space_before=6)
para(
    "训练平滑收敛，约 2200 个 epoch 时早停，最佳验证 MSE 约 3.5×10⁻⁵（图 4）。在留出的 500 个"
    "测试样本上，RMSE = 0.0060，MAE = 0.0042。代表性测试谱（图 5）显示，MLP 在大部分波段都能"
    "复现 TMM 参考谱的峰位和整体线型，只在尖锐极值附近有小偏差。")
fig("fig4_loss.png", "图 4. MLP 代理模型的训练与验证 MSE 损失（纵轴对数）。")
fig("fig5_spectra.png", "图 5. 三个代表性测试样本：TMM 参考（黑实线）与 MLP 预测（红虚线）对比；绿色点线为 480 nm。",
    width=15.5)

h("3.2 训练集规模的影响", 11, space_before=6)
para(
    "测试 RMSE 从 500 样本的 0.0262，降到 1000 的 0.0135、2000 的 0.0097、4000 的 0.0060"
    "（图 6）。数据翻倍时误差先近似减半，但随着样本增多，增量收益不断缩小，呈现边际递减"
    "而非线性改善。")
fig("fig6_ablation.png", "图 6. 训练集规模对测试 RMSE 的影响；验证集与测试集固定不变。")

h("3.3 MLP 辅助薄膜设计", 11, space_before=6)
para(
    "用 MLP 筛选 10000 个新候选，再对前 10 名做 TMM 复核，得到表 1。最优设计 "
    "d = (84.2, 125.4, 63.5, 155.2) nm，TMM 复核的 R(480 nm) 仅 0.0003，即在目标波长实现了"
    "近乎完全的相消干涉（图 7）；该设计的 MLP 预测谱与 TMM 复核谱几乎重合。")

tbl = doc.add_table(rows=6, cols=7); tbl.style = "Light Grid Accent 1"
hdr = ["排名", "d1/nm", "d2/nm", "d3/nm", "d4/nm", "MLP R@480", "TMM R@480"]
data = [
    ["1", "84.2", "125.4", "63.5", "155.2", "-0.0060", "0.0003"],
    ["2", "85.5", "120.4", "67.1", "40.7", "-0.0096", "0.0012"],
    ["3", "117.1", "70.7", "128.5", "58.3", "-0.0044", "0.0013"],
    ["4", "122.4", "56.5", "128.9", "87.0", "-0.0061", "0.0025"],
    ["5", "123.3", "176.6", "77.8", "90.4", "-0.0039", "0.0045"],
]
for j, txt in enumerate(hdr):
    c = tbl.rows[0].cells[j]; c.text = ""
    set_cn(c.paragraphs[0].add_run(txt), size=9, bold=True)
for i, row in enumerate(data, start=1):
    for j, txt in enumerate(row):
        c = tbl.rows[i].cells[j]; c.text = ""
        set_cn(c.paragraphs[0].add_run(txt), size=9)
cap = doc.add_paragraph(); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_cn(cap.add_run("表 1. MLP 筛选并经 TMM 复核的 Top5 减反设计（λ = 480 nm）。"), size=9, italic=True)

fig("fig7_design.png", "图 7. MLP 选出的最优设计：TMM 复核（黑）与 MLP 预测（红虚线）光谱，在 480 nm 处反射率降至 R≈0。",
    width=12)

h("3.4 代表性失败案例", 11, space_before=6)
para(
    "误差最大的测试样本（单样本 MSE ≈ 8.6×10⁻⁴，图 8）中，MLP 系统性地把两个反射率波峰估高，"
    "并略微移动了峰位。绝对误差不大，但正好暴露了代理模型的软肋：窄而高的峰被平滑回归“平均”"
    "得不够准。")
fig("fig8_failure.png", "图 8. 代表性失败案例：MLP 高估并轻微位移了测试谱中的反射率峰。", width=12)

# ---------------- 4 讨论 ----------------
h("4. 讨论")
para(
    "Q1：MLP 以约 0.006 的 RMSE 预测 41 点光谱，相对于 0–1 的反射率范围已经很小，足以区分膜系"
    "设计的优劣。误差集中在峰位和曲线交叉处——那里膜厚变化引起的光谱变化最剧烈，符合平滑回归"
    "在窄特征附近“平均化”的特点。")
para(
    "Q2：增加训练数据可靠地降低误差，但边际递减明显。500→1000 误差约减半，2000→4000 只带来"
    "有限改善。对这个四层问题，几千次 TMM 仿真就足以训练出可用的代理模型，再多数据收益递减。")
para(
    "Q3：代理模型的价值在筛选。MLP 排序 10000 个候选只需毫秒级，而 TMM 逐条计算要慢得多；最终"
    "最优设计 R(480 nm) = 0.0003。值得注意的是，MLP 偶尔会输出略负的反射率（如 −0.006），这是"
    "回归模型没有 [0,1] 约束导致的、物理上不可能的值——这正是最终候选必须用 TMM 重算并重排的"
    "原因。两者分工：MLP 负责广度，TMM 负责可信。")
para(
    "局限：本模型固定了材料、层数和正入射条件，只是 [40,180] nm 训练范围内的插值器，外推能力"
    "有限；网络本身没有物理约束，原始输出不能直接当作精确结果。后续可加入有界输出、考虑色散，"
    "或采用预测与优化更耦合的串联模型。")

# ---------------- 5 结论 ----------------
h("5. 结论")
para(
    "一个紧凑的 MLP（4–128–128–64–41）能准确学习四层介质膜的 TMM 膜厚—光谱映射，测试 RMSE = "
    "0.0060（Q1）。预测误差随训练数据增加而单调下降，但边际递减（500 样本 0.026 → 4000 样本 "
    "0.006）（Q2）。作为快速代理，MLP 筛选 10000 个候选，经 TMM 复核后在 480 nm 找到 R = 0.0003 "
    "的减反设计（Q3）。整个“物理仿真—数据代理—物理校验”流程，构成了一个虽小但完整的科学智能"
    "设计闭环。")

# ---------------- 数据可用性 ----------------
h("数据与代码可用性")
para(
    "数据集由本研究的 TMM 代码生成。全部源码、训练与筛选脚本、依赖文件及复现说明见 GitHub 仓库"
    "（链接待替换；λtarget = 480 nm，seed = 270034，design_seed = 270035）。")

# ---------------- 参考文献 ----------------
h("参考文献")
refs = [
    "Ma, T.; Ma, M.; Guo, L.J. Optical multilayer thin film structure inverse design: From optimization to deep learning. iScience 2025, 28, 112222.",
    "Macleod, H.A. Thin-Film Optical Filters, 5th ed.; CRC Press: Boca Raton, 2017.",
    "Born, M.; Wolf, E. Principles of Optics, 7th ed.; Cambridge University Press: Cambridge, 1999.",
    "Rumelhart, D.E.; Hinton, G.E.; Williams, R.J. Learning representations by back-propagating errors. Nature 1986, 323, 533–536.",
    "Goodfellow, I.; Bengio, Y.; Courville, A. Deep Learning; MIT Press: Cambridge, 2016.",
    "Paszke, A. et al. PyTorch: An imperative style, high-performance deep learning library. Adv. Neural Inf. Process. Syst. 2019, 32, 8024–8035.",
    "Piegari, A.; Flory, F. (Eds.) Optical Thin Films and Coatings: From Materials to Applications; Woodhead Publishing: Cambridge, 2018.",
    "Goldberg, D.E. Genetic Algorithms in Search, Optimization, and Machine Learning; Addison-Wesley: Reading, 1989.",
    "Molesky, S. et al. Machine learning and nanophotonic design. Nat. Photonics 2018, 12, 659–670.",
]
for i, ref in enumerate(refs, 1):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    set_cn(p.add_run(f"[{i}] {ref}"), size=9)

doc.save(OUT)
print("saved", OUT)
